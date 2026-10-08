# 2단계: df.pkl로 지역 현황 엑셀(요약·연도별 개설·매도 후보·데이터) 생성
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
df=pd.read_pickle('df.pkl'); s=df[df.region.notna()].copy()
s['후보']=((s['학원교습소명']=='학원')&s['분야명'].isin(['입시.검정 및 보습','국제화'])&(s.age>=5)&s.cap.between(20,80)).map({True:'Y',False:'N'})
order=['목동','상계·중계','대치','분당 정자','분당 수내','분당 서현']
s['ord']=s.region.map({k:i for i,k in enumerate(order)})
s=s.sort_values(['ord','age'],ascending=[True,False])
F='Arial'; hdr=Font(name=F,bold=True,color='FFFFFF'); hfill=PatternFill('solid',fgColor='22463A')
base=Font(name=F); note=Font(name=F,italic=True,color='5F6B65',size=9); title=Font(name=F,bold=True,size=14)
thin=Border(bottom=Side(style='thin',color='D9DED6'))
wb=Workbook()
def header(ws,row,cols):
    for i,c in enumerate(cols,1):
        x=ws.cell(row,i,c); x.font=hdr; x.fill=hfill; x.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
# 데이터 sheet
d=wb.active; d.title='데이터'
cols=['지역','구분','학원명','분야','교습과정','개설일','업력(년)','개설연도','일시수용인원','매도후보','도로명주소','상세주소','전화번호']
header(d,1,cols)
for r,(_,x) in enumerate(s.iterrows(),2):
    vals=[x.region,x['학원교습소명'],x['학원명'],x['분야명'],x['교습과정명'],x['개설일자'],round(float(x.age),1) if pd.notna(x.age) else None,
          int(x.est_year) if pd.notna(x.est_year) and x.est_year<2100 else None,int(x.cap) if pd.notna(x.cap) else None,x['후보'],
          x['도로명주소'],(x['도로명상세주소'] or '').strip(' ,'),None if x['전화번호'] in (None,'null') else x['전화번호']]
    for c,v in enumerate(vals,1): d.cell(r,c,v).font=base
N=len(s)+1
for i,w in enumerate([10,8,34,16,14,10,9,9,11,9,40,30,14],1): d.column_dimensions[get_column_letter(i)].width=w
d.freeze_panes='A2'; d.auto_filter.ref=f'A1:M{N}'
rng=lambda col:f"데이터!${col}$2:${col}${N}"
A,B,D,G,H,I,J=[rng(c) for c in 'ABDGHIJ']
# 요약
w=wb.create_sheet('요약',0)
w['A1']='바톤 시작 지역 학원 현황'; w['A1'].font=title
w['A2']='출처: NEIS 학원교습소정보 2026년 3월 (운영 중인 학원·교습소만 수록). 지역은 상세주소의 법정동으로 분류.'; w['A2'].font=note
cols=['지역','전체','학원','교습소','교습소 비중','입시·보습 비중','평균 업력(년)','업력 10년 이상','업력 3년 미만','매도 후보']
header(w,4,cols)
for i,k in enumerate(order,5):
    w.cell(i,1,k)
    w.cell(i,2,f'=COUNTIFS({A},$A{i})')
    w.cell(i,3,f'=COUNTIFS({A},$A{i},{B},"학원")')
    w.cell(i,4,f'=COUNTIFS({A},$A{i},{B},"교습소")')
    w.cell(i,5,f'=IF(B{i}=0,0,D{i}/B{i})')
    w.cell(i,6,f'=IF(B{i}=0,0,COUNTIFS({A},$A{i},{D},"입시.검정 및 보습")/B{i})')
    w.cell(i,7,f'=IFERROR(AVERAGEIFS({G},{A},$A{i}),0)')
    w.cell(i,8,f'=IF(B{i}=0,0,COUNTIFS({A},$A{i},{G},">=10")/B{i})')
    w.cell(i,9,f'=IF(B{i}=0,0,COUNTIFS({A},$A{i},{G},"<3")/B{i})')
    w.cell(i,10,f'=COUNTIFS({A},$A{i},{J},"Y")')
t=5+len(order)
w.cell(t,1,'합계').font=Font(name=F,bold=True)
for c in (2,3,4,10): w.cell(t,c,f'=SUM({get_column_letter(c)}5:{get_column_letter(c)}{t-1})').font=Font(name=F,bold=True)
for r in range(5,t+1):
    for c in range(1,11):
        x=w.cell(r,c); x.border=thin
        if x.font!=Font(name=F,bold=True): x.font=base
        if c in (5,6,8,9): x.number_format='0.0%'
        if c==7: x.number_format='0.0'
        if c in (2,3,4,10): x.number_format='#,##0'
w.cell(t+2,1,'매도 후보 기준').font=Font(name=F,bold=True)
for j,txt in enumerate(['학원(교습소 제외), 분야 = 입시·검정 및 보습 또는 국제화(영어)',
    '개설 5년 이상 (2026-03-31 기준)','일시수용인원 20~80명 (중소형)',
    '가설 기준이며 실제 매도 의향과는 무관. 인터뷰 대상 우선순위용'],t+3):
    w.cell(j,1,'· '+txt).font=base
w.cell(t+8,1,'참고').font=Font(name=F,bold=True)
for j,txt in enumerate(['상세주소에서 동을 읽지 못한 행(전국 약 12.5%)은 빠져 있어 실제 수보다 조금 적을 수 있음',
    '원본에는 운영 중인 곳만 있어 폐원 수는 알 수 없음. 매달 파일을 비교하면 폐원·신규를 잡을 수 있음',
    '인당수강료 항목은 원본에서 비어 있음'],t+9):
    w.cell(j,1,'· '+txt).font=base
w.column_dimensions['A'].width=14
for c in range(2,11): w.column_dimensions[get_column_letter(c)].width=12
w.row_dimensions[4].height=32
# 연도별 개설
y=wb.create_sheet('연도별 개설',1)
y['A1']='연도별 개설 수 (현재 운영 중인 곳 기준)'; y['A1'].font=title
y['A2']='과거 연도일수록 이미 문 닫은 곳이 빠져 적게 보임(생존자만 집계).'; y['A2'].font=note
years=list(range(2016,2026))
header(y,4,['지역']+[str(v) for v in years])
for i,k in enumerate(order,5):
    y.cell(i,1,k).font=base
    for j,v in enumerate(years,2):
        c=y.cell(i,j,f'=COUNTIFS({A},$A{i},{H},{v})'); c.font=base; c.number_format='#,##0'
y.column_dimensions['A'].width=14
# 매도 후보 sheet (static list)
c=wb.create_sheet('매도 후보',2)
cs=s[s['후보']=='Y']
cols2=['지역','학원명','교습과정','개설일','업력(년)','일시수용인원','도로명주소','상세주소','전화번호','연락 상태','메모']
header(c,1,cols2)
for r,(_,x) in enumerate(cs.iterrows(),2):
    vals=[x.region,x['학원명'],x['교습과정명'],x['개설일자'],round(float(x.age),1),int(x.cap),x['도로명주소'],(x['도로명상세주소'] or '').strip(' ,'),
          None if x['전화번호'] in (None,'null') else x['전화번호'],None,None]
    for k,v in enumerate(vals,1): c.cell(r,k,v).font=base
yellow=PatternFill('solid',fgColor='FFF4C2')
for r in range(2,len(cs)+2):
    c.cell(r,10).fill=yellow; c.cell(r,11).fill=yellow
for i,wd in enumerate([10,34,12,10,9,11,38,28,14,12,30],1): c.column_dimensions[get_column_letter(i)].width=wd
c.freeze_panes='A2'; c.auto_filter.ref=f'A1:K{len(cs)+1}'
c.cell(len(cs)+3,1,'노란 칸(연락 상태, 메모)은 인터뷰 진행하며 직접 채우세요. 예: 연락 상태 = 미연락/통화/인터뷰 완료/매도 의향 있음').font=note
wb.save('baton_regions_2026-03.xlsx'); print(len(cs), N)
