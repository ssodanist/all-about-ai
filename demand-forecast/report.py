"""엑셀 리포트 작성. 발주 시트는 수식으로 연결되어 있어 현재고·리드타임 등을 고치면 발주량이 다시 계산됩니다."""
from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=14, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_INPUT = Font(name=FONT, size=10, color="0000FF")          # 파란 글씨 = 직접 수정 가능한 입력값
F_LINK = Font(name=FONT, size=10, color="008000")           # 초록 글씨 = 다른 시트 참조
F_NOTE = Font(name=FONT, size=9, italic=True, color="666666")
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_KEY = PatternFill("solid", fgColor="FFFF00")           # 노란 칸 = 매일 채워야 하는 값
FILL_ORDER = PatternFill("solid", fgColor="E2EFDA")
FILL_HOL = PatternFill("solid", fgColor="FCE4D6")
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))
WEEKDAY_KR = "월화수목금토일"
HEAD_ROW, FIRST = 4, 5


def _header(ws, cols, widths=None):
    for j, name in enumerate(cols, 1):
        c = ws.cell(HEAD_ROW, j, name)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(name, 11)
    ws.row_dimensions[HEAD_ROW].height = 30
    ws.freeze_panes = ws.cell(FIRST, 4)


def _title(ws, title, note):
    ws["A1"] = title; ws["A1"].font = F_TITLE
    ws["A2"] = note; ws["A2"].font = F_NOTE


def write_excel(path, plan, daily, future, accuracy, effects, holidays, dates, cfg, H, model_desc):
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    n = len(plan); last = FIRST + n - 1
    hol_set = set(pd.DatetimeIndex(holidays).normalize())

    # ---------------- 일별예측 ----------------
    wd = wb.active; wd.title = "일별예측"
    _title(wd, "점포×품목 일별 수요예측 (단위: 개)",
           "주황 열 = 공휴일. 값은 모델 산출 결과이며, 발주추천 시트가 이 값을 합산해 사용합니다.")
    heads = ["점포", "SKU", "품목명"] + [f"{d:%m/%d}({WEEKDAY_KR[d.dayofweek]})" for d in future]
    _header(wd, heads, {"품목명": 18})
    for j, d in enumerate(future, 4):
        wd.column_dimensions[get_column_letter(j)].width = 9
        if d.normalize() in hol_set:
            wd.cell(HEAD_ROW, j).fill = PatternFill("solid", fgColor="C65911")
    for i, row in enumerate(daily.itertuples(index=False), FIRST):
        for j, v in enumerate(row, 1):
            c = wd.cell(i, j, v); c.font = F_BASE
            if j > 3:
                c.number_format = "0.0"
                if future[j - 4].normalize() in hol_set:
                    c.fill = FILL_HOL
    dl, dr = "D", get_column_letter(3 + H)
    wd.auto_filter.ref = f"A{HEAD_ROW}:{dr}{last}"

    # ---------------- 발주추천 ----------------
    wo = wb.create_sheet("발주추천", 0)
    _title(wo, f"발주 추천 ({future[0]:%Y-%m-%d} 발주 기준)",
           "노란 칸(현재고·입고예정)을 매일 입력하세요. 파란 글씨(리드타임·서비스수준 등)는 수정 가능하며 "
           "발주수량은 자동 재계산됩니다.")
    cols = ["점포", "SKU", "품목명", "ABC", "수요패턴", "선택모델", "리드타임(일)", "발주주기(일)", "보호기간(일)",
            "보호기간 예측수요", "일환산 예측오차σ", "서비스수준", "안전재고", "유통기한(일)", "유통기한 상한",
            "목표재고", "현재고", "입고예정", "필요량", "박스입수", "MOQ", "발주수량", "단가(원)", "발주금액(원)",
            "발주후 재고일수"]
    _header(wo, cols, {"품목명": 18, "수요패턴": 10, "선택모델": 12, "발주금액(원)": 13})
    notes = {
        "보호기간(일)": "리드타임 + 발주주기. 이번 발주로 다음 입고 전까지 버텨야 하는 기간.",
        "보호기간 예측수요": "일별예측 시트에서 보호기간 일수만큼 합산.",
        "일환산 예측오차σ": "백테스트에서 측정한 보호기간 누적 예측오차(RMSE)를 하루 단위로 환산한 값.",
        "안전재고": "NORMSINV(서비스수준) × σ × √보호기간",
        "유통기한 상한": "리드타임+유통기한 동안 팔 수 있는 양. 이를 넘게 쌓으면 폐기 위험. '-'는 제한 없음.",
        "서비스수준": "보호기간 중 결품이 나지 않을 확률 목표 (ABC 등급별 기본값).",
        "발주수량": "필요량을 박스입수 단위로 올림, MOQ 이상.",
    }
    for j, name in enumerate(cols, 1):
        if name in notes:
            wo.cell(HEAD_ROW, j).comment = Comment(notes[name], "forecast")

    rng = lambda r, op: (f"SUMPRODUCT('일별예측'!${dl}{r}:${dr}{r}*((COLUMN('일별예측'!${dl}{r}:${dr}{r})"
                         f"-COLUMN('일별예측'!${dl}{r})+1)<={op}))")
    for i, p in enumerate(plan.itertuples(index=False), FIRST):
        r = i
        vals = [p.store, p.sku, p.name, p.abc, p.pattern, p.model, p.lead_time, p.review_period,
                f"=G{r}+H{r}", "=" + rng(r, f"I{r}"), p.sigma_day, p.service_level,
                f"=NORMSINV(L{r})*K{r}*SQRT(I{r})", p.shelf_life,
                f'=IF(G{r}+N{r}>{H},"-",{rng(r, f"(G{r}+N{r})")})',
                f"=IF(ISNUMBER(O{r}),MIN(J{r}+M{r},O{r}),J{r}+M{r})",
                p.on_hand, p.on_order, f"=MAX(0,P{r}-Q{r}-R{r})", p.case_pack, p.moq,
                f"=IF(S{r}<=0,0,MAX(U{r},ROUNDUP(S{r}/MAX(T{r},1),0)*MAX(T{r},1)))",
                p.unit_cost, f"=V{r}*W{r}", f'=IFERROR((Q{r}+R{r}+V{r})/(J{r}/I{r}),"-")']
        for j, v in enumerate(vals, 1):
            c = wo.cell(r, j, v); c.font = F_BASE; c.border = THIN
        for col in "GHLNTUW":
            wo[f"{col}{r}"].font = F_INPUT
        for col in "QR":
            wo[f"{col}{r}"].font = F_INPUT; wo[f"{col}{r}"].fill = FILL_KEY
        wo[f"J{r}"].font = F_LINK; wo[f"O{r}"].font = F_LINK
        wo[f"V{r}"].font = F_BOLD; wo[f"V{r}"].fill = FILL_ORDER
        for col, fmt in {"J": "#,##0.0", "K": "0.00", "L": "0.0%", "M": "#,##0.0", "O": "#,##0.0",
                         "P": "#,##0.0", "S": "#,##0.0", "V": "#,##0;-#,##0;-", "W": "#,##0",
                         "X": "#,##0;-#,##0;-", "Y": "0.0"}.items():
            wo[f"{col}{r}"].number_format = fmt
    wo.auto_filter.ref = f"A{HEAD_ROW}:Y{last}"

    # ---------------- 예측정확도 ----------------
    wa = wb.create_sheet("예측정확도")
    bt_from = dates[-1] - pd.Timedelta(days=cfg.backtest_days + H - 1)
    _title(wa, "백테스트 예측정확도 (과거 시점에서 실제로 예측해 보고 맞춰본 결과)",
           f"최근 약 {cfg.backtest_days}일 동안 매일을 예측 시점으로 가정, 1~{H}일 앞 예측을 실제와 비교 "
           f"(검증구간 시작 ≈ {bt_from:%Y-%m-%d}). 결품일은 제외. 벤치마크 = 지난주 같은 요일 판매량 그대로 쓰기.")
    acols = ["점포", "SKU", "품목명", "수요패턴", "선택모델", "실제판매 합계", "절대오차 합계", "편향(+과대)",
             "WAPE(오차율)", "벤치마크 절대오차", "벤치마크 WAPE", "오차 감소율"]
    _header(wa, acols, {"품목명": 18, "선택모델": 12})
    for i, a in enumerate(accuracy.itertuples(index=False), FIRST):
        vals = [a.store, a.sku, a.name, a.pattern, a.model, a.actual, a.abs_err, a.bias,
                f'=IFERROR(G{i}/F{i},"-")', a.bench_abs_err, f'=IFERROR(J{i}/F{i},"-")',
                f'=IFERROR(1-G{i}/J{i},"-")']
        for j, v in enumerate(vals, 1):
            c = wa.cell(i, j, v); c.font = F_BASE; c.border = THIN
        for col in "FGHJ":
            wa[f"{col}{i}"].number_format = "#,##0;-#,##0;-"
        for col in "IKL":
            wa[f"{col}{i}"].number_format = "0.0%"
    alast = FIRST + len(accuracy) - 1
    wa.auto_filter.ref = f"A{HEAD_ROW}:L{alast}"

    # ---------------- 특수일효과 ----------------
    we = wb.create_sheet("특수일효과")
    _title(we, "학습된 특수일 효과 (1.00 = 평소와 같음, 1.20 = 20% 증가)",
           "과거 실적 / 평소 기준선 비율. 관측일이 적으면 1.00 쪽으로 보수적으로 수축됨.")
    for j, h in enumerate(["점포", "공휴일 승수", "연휴 전날 승수"], 1):
        c = we.cell(HEAD_ROW, j, h); c.font, c.fill = F_HEAD, FILL_HEAD
    for i, st in enumerate(effects["holiday"].index, FIRST):
        we.cell(i, 1, st).font = F_BASE
        for j, key in ((2, "holiday"), (3, "pre_holiday")):
            c = we.cell(i, j, float(effects[key][st])); c.number_format = "0.00"; c.font = F_BASE
    for j, h in enumerate(["SKU", "프로모션 승수"], 5):
        c = we.cell(HEAD_ROW, j, h); c.font, c.fill = F_HEAD, FILL_HEAD
    for i, sku in enumerate(effects["promo"].index, FIRST):
        we.cell(i, 5, sku).font = F_BASE
        c = we.cell(i, 6, float(effects["promo"][sku])); c.number_format = "0.00"; c.font = F_BASE
    for col, w in zip("ABCDEF", (10, 12, 14, 3, 12, 14)):
        we.column_dimensions[col].width = w

    # ---------------- 요약 ----------------
    ws = wb.create_sheet("요약", 0)
    ws.column_dimensions["A"].width = 26; ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 60
    ws["A1"] = "수요예측 · 발주추천 요약"; ws["A1"].font = F_TITLE
    ws["A2"] = f"데이터 {dates[0]:%Y-%m-%d} ~ {dates[-1]:%Y-%m-%d} · 예측 {future[0]:%Y-%m-%d}부터 {H}일"
    ws["A2"].font = F_NOTE
    O, A = "'발주추천'!", "'예측정확도'!"
    rows = [
        ("총 발주금액(원)", f"=SUM({O}X{FIRST}:X{last})", "#,##0", "발주추천 시트 합계"),
        ("총 발주수량(개)", f"=SUM({O}V{FIRST}:V{last})", "#,##0", ""),
        ("발주 필요 건수", f'=COUNTIF({O}V{FIRST}:V{last},">0")', "#,##0", f"전체 {n}개 점포×품목 중"),
        ("총 안전재고(개)", f"=SUM({O}M{FIRST}:M{last})", "#,##0", "예측오차를 흡수하는 완충 재고"),
        ("", None, None, ""),
        ("예측 오차율 WAPE – 이 엔진", f"=IFERROR(SUM({A}G{FIRST}:G{alast})/SUM({A}F{FIRST}:F{alast}),\"-\")",
         "0.0%", "Σ|예측−실제| / Σ실제. 낮을수록 좋음"),
        ("예측 오차율 WAPE – 지난주 같은요일", f"=IFERROR(SUM({A}J{FIRST}:J{alast})/SUM({A}F{FIRST}:F{alast}),\"-\")",
         "0.0%", "현업에서 흔히 쓰는 감(感) 기반 방식의 대용치"),
        ("오차 감소율", "=IFERROR(1-B10/B11,\"-\")", "0.0%", "벤치마크 대비 얼마나 덜 틀리는가"),
        ("예측 편향", f"=IFERROR(SUM({A}H{FIRST}:H{alast})/SUM({A}F{FIRST}:F{alast}),\"-\")",
         "+0.0%;-0.0%;0.0%", "+면 과대(폐기 위험), −면 과소(결품 위험). ±3% 이내가 양호"),
    ]
    for i, (label, f, fmt, note) in enumerate(rows, 5):
        ws.cell(i, 1, label).font = F_BOLD
        if f:
            c = ws.cell(i, 2, f); c.number_format = fmt; c.font = F_LINK
        ws.cell(i, 3, note).font = F_NOTE
    r0 = 5 + len(rows) + 1
    ws.cell(r0, 1, "점포별 발주금액").font = F_BOLD
    for j, h in enumerate(["점포", "발주금액(원)", "발주 건수"], 1):
        c = ws.cell(r0 + 1, j, h); c.font, c.fill = F_HEAD, FILL_HEAD
    for i, st in enumerate(sorted(plan["store"].unique()), r0 + 2):
        ws.cell(i, 1, st).font = F_BASE
        c = ws.cell(i, 2, f'=SUMIFS({O}X{FIRST}:X{last},{O}A{FIRST}:A{last},A{i})')
        c.number_format = "#,##0"; c.font = F_LINK
        c = ws.cell(i, 3, f'=COUNTIFS({O}A{FIRST}:A{last},A{i},{O}V{FIRST}:V{last},">0")')
        c.font = F_LINK

    # ---------------- 사용법·가정 ----------------
    wg = wb.create_sheet("사용법·가정")
    wg.column_dimensions["A"].width = 120
    sl = cfg.service_level
    lines = [
        ("사용법", True),
        ("1) '발주추천' 시트의 노란 칸(현재고, 입고예정)을 입력하면 발주수량이 자동 계산됩니다.", False),
        ("2) 파란 글씨는 조정 가능한 가정값입니다(리드타임, 발주주기, 서비스수준, 유통기한, 박스입수, MOQ, 단가).", False),
        ("3) 초록 글씨는 다른 시트를 참조하는 수식, 검정 글씨는 계산 수식입니다.", False),
        ("", False),
        ("가정 및 출처", True),
        ("· 판매·결품·프로모션 이력, 품목 마스터(리드타임·박스입수 등), 재고는 사용자가 제공한 CSV 값입니다.", False),
        (f"· 서비스수준 기본값: A등급 {sl['A']:.0%}, B등급 {sl['B']:.0%}, C등급 {sl['C']:.0%} "
         "(최근 13주 매출금액 누적 80%/95% 기준 ABC 분류). 업계 일반 관행 값이며 회사 정책에 맞게 조정하세요.", False),
        ("· 결품일 판매량은 실제 수요보다 작으므로 학습·검증에서 제외했습니다.", False),
        ("· 공휴일은 한국 법정공휴일·대체공휴일(2024~2027) 내장 목록 기준. 임시공휴일은 holidays.csv로 추가하세요.", False),
        ("· 안전재고 = NORMSINV(서비스수준) × (백테스트 보호기간 누적오차 σ) — 오차를 정규분포로 근사.", False),
        ("· 목표재고는 '리드타임+유통기한' 동안의 예측수요를 넘지 않도록 상한을 둡니다(폐기 방지).", False),
        ("· 일별예측 값은 모델 결과(하드코딩)이며 예측일수 밖의 보호기간은 계산할 수 없습니다.", False),
        ("", False),
        ("모델 설명", True),
    ]
    lines += [(f"· {k}: {v}", False) for k, v in model_desc.items()]
    lines += [("· 시계열(점포×품목)마다 백테스트 점수(MAE+|편향|)가 가장 좋은 모델을 자동 선택합니다.", False)]
    for i, (t, bold) in enumerate(lines, 1):
        c = wg.cell(i, 1, t); c.font = F_BOLD if bold else F_BASE

    wb.save(path)
